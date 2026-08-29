"""Writing laid-out pages into a .pptx file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .layout import Line, Page
from .theme import Theme

ALIGNMENTS = {
    "center": PP_ALIGN.CENTER,
    "left": PP_ALIGN.LEFT,
    "right": PP_ALIGN.RIGHT,
}


def _rgb(hex_string: str) -> RGBColor:
    return RGBColor.from_string(hex_string.lstrip("#").upper())


def render(pages: list[Page], theme: Theme, out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(theme.width_in)
    prs.slide_height = Inches(theme.height_in)
    blank_layout = prs.slide_layouts[6]

    for page in pages:
        slide = prs.slides.add_slide(blank_layout)
        _paint_background(slide, theme, page)

        if page.kind == "image":
            _draw_image(slide, theme, page, prs)
        elif page.kind != "blank":
            _draw_content(slide, theme, page)

        if page.notes:
            slide.notes_slide.notes_text_frame.text = page.notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def _paint_background(slide, theme: Theme, page: Page) -> None:
    name = "title_background" if page.kind in {"title", "blank"} else "background"
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(theme.color(name))


def _textbox(slide, *, left, top, width, height, anchor=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def _style(run, *, font: str, size: float, color: str, bold: bool = False,
           italic: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)


def _draw_content(slide, theme: Theme, page: Page) -> None:
    margin_x = Inches(float(theme.data["margins"]["x_in"]))
    margin_top = Inches(float(theme.data["margins"]["top_in"]))
    margin_bottom = Inches(float(theme.data["margins"]["bottom_in"]))
    width = Inches(theme.body_width_in)
    align = ALIGNMENTS.get(str(page.meta.get("align", "center")), PP_ALIGN.CENTER)

    body_top = margin_top
    body_bottom = Inches(theme.height_in) - margin_bottom

    if page.header:
        header_h = Inches(theme.pt("heading_pt") * 1.6 / 72.0)
        frame = _textbox(
            slide,
            left=margin_x,
            top=margin_top,
            width=width,
            height=header_h,
            anchor=MSO_ANCHOR.TOP,
        )
        para = frame.paragraphs[0]
        para.alignment = align
        _style(
            para.add_run(),
            font=theme.font("ui"),
            size=theme.pt("heading_pt"),
            color=theme.color("accent"),
        )
        para.runs[0].text = page.header
        body_top = margin_top + header_h + Inches(float(theme.opt("header_gap_in")))

    if page.footer:
        footer_h = Inches(theme.pt("footer_pt") * 2.2 / 72.0)
        frame = _textbox(
            slide,
            left=margin_x,
            top=Inches(theme.height_in) - margin_bottom * 0.5 - footer_h,
            width=width,
            height=footer_h,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        para = frame.paragraphs[0]
        para.alignment = align
        run = para.add_run()
        run.text = page.footer
        _style(
            run,
            font=theme.font("ui"),
            size=theme.pt("footer_pt"),
            color=theme.color("muted"),
        )
        body_bottom -= footer_h

    if theme.opt("show_page_numbers") and page.total > 1:
        _draw_page_number(slide, theme, page)

    title = str(page.meta.get("title") or "")
    lines = page.lines
    if not title and not lines:
        return

    # Title and body share one middle-anchored frame so they read as a single
    # block. Two separate boxes leave a hole between the heading and the text.
    height = max(Inches(0.5), body_bottom - body_top)
    frame = _textbox(slide, left=margin_x, top=body_top, width=width, height=height)
    first = True

    if title:
        size = float(
            theme.pt("title_pt")
            if page.kind == "title"
            else page.meta.get("title_pt") or theme.pt("card_title_pt")
        )
        para = frame.paragraphs[0]
        para.alignment = align
        para.line_spacing = 1.1
        run = para.add_run()
        run.text = title
        _style(
            run,
            font=theme.font("heading"),
            size=size,
            color=theme.color("text"),
            bold=True,
        )
        first = False
        if lines:
            spacer = frame.add_paragraph()
            spacer.line_spacing = 1.0
            gap = spacer.add_run()
            gap.text = ""
            _style(
                gap,
                font=theme.font("body"),
                size=size * 0.5,
                color=theme.color("background"),
            )

    if lines:
        _fill_lines(frame, theme, page, lines, align, first_paragraph=first)


def _indent(para, inches: float, *, hanging: bool) -> None:
    """Align liturgy text in one column to the right of the speaker labels.

    A labelled line hangs its "L:" back into the gutter; an unlabelled
    continuation line simply starts at the gutter, so every line of an exchange
    lines up regardless of which ones carry a label.
    """
    emu = int(Inches(inches))
    properties = para._p.get_or_add_pPr()
    properties.set("marL", str(emu))
    properties.set("indent", str(-emu if hanging else 0))


def _fill_lines(
    frame, theme: Theme, page: Page, lines: list[Line], align, first_paragraph=True
) -> None:
    spacing = float(theme.opt("line_spacing"))
    show_labels = bool(theme.opt("show_labels"))

    # Reserve a gutter wide enough for the longest speaker label on this slide.
    widest = max((len(line.label) for line in lines if line.label), default=0)
    gutter_in = ((widest + 2) * page.font_pt * 0.8 * 0.5) / 72.0 if widest else 0.0

    for i, line in enumerate(lines):
        if first_paragraph and i == 0:
            para = frame.paragraphs[0]
        else:
            para = frame.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        if gutter_in and align == PP_ALIGN.LEFT:
            _indent(para, gutter_in, hanging=bool(line.label and show_labels))

        if line.role == "blank":
            run = para.add_run()
            run.text = ""
            _style(
                run,
                font=theme.font("body"),
                size=page.font_pt * 0.45,
                color=theme.color("background"),
            )
            continue

        if line.label and show_labels:
            label_run = para.add_run()
            label_run.text = f"{line.label}:  "
            _style(
                label_run,
                font=theme.font("ui"),
                size=page.font_pt * 0.8,
                color=theme.color("accent"),
                bold=True,
            )

        run = para.add_run()
        run.text = line.text
        _style(
            run,
            font=theme.font("body"),
            size=page.font_pt,
            color=theme.color("text"),
            bold=(line.role == "people"),
        )


def _draw_page_number(slide, theme: Theme, page: Page) -> None:
    size = theme.pt("footer_pt")
    box_w = Inches(1.4)
    frame = _textbox(
        slide,
        left=Inches(theme.width_in) - box_w - Inches(0.35),
        top=Inches(theme.height_in) - Inches(0.5),
        width=box_w,
        height=Inches(0.32),
        anchor=MSO_ANCHOR.BOTTOM,
    )
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    run = para.add_run()
    run.text = f"{page.index} / {page.total}"
    _style(run, font=theme.font("ui"), size=size, color=theme.color("muted"))


def _draw_image(slide, theme: Theme, page: Page, prs) -> None:
    from PIL import Image

    path = page.meta["path"]
    slide_w, slide_h = prs.slide_width, prs.slide_height
    with Image.open(path) as img:
        img_w, img_h = img.size

    # Cover the slide, cropping the overflow rather than letterboxing.
    scale = max(slide_w / img_w, slide_h / img_h)
    width, height = int(img_w * scale), int(img_h * scale)
    slide.shapes.add_picture(
        path,
        int((slide_w - width) / 2),
        int((slide_h - height) / 2),
        width=width,
        height=height,
    )

    caption = page.meta.get("caption")
    if caption:
        frame = _textbox(
            slide,
            left=Inches(float(theme.data["margins"]["x_in"])),
            top=slide_h - Inches(1.3),
            width=Inches(theme.body_width_in),
            height=Inches(0.9),
            anchor=MSO_ANCHOR.BOTTOM,
        )
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = caption
        _style(
            run,
            font=theme.font("heading"),
            size=theme.pt("subtitle_pt"),
            color=theme.color("text"),
            bold=True,
        )
